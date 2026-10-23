from pathlib import Path
import hashlib
import json
import re
import shutil
import zipfile
import fitz
import openpyxl

root = Path(__file__).parent
old = Path('/mnt/c/Users/sg564/Desktop/ICA3PP_237_CameraReady_20261008_no_fig8')
src, out = root / 'source', root / 'delivery'
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
original_hashes = {p.name: sha(p) for p in old.iterdir() if p.is_file()}
changed = {'main.tex', 'figures/fig_eval_timeline_overall.pdf', 'figures/fig_eval_timeline_groups.pdf'}
with zipfile.ZipFile(old / '237_Source.zip') as z:
    with zipfile.ZipFile(out / '237_Source.zip', 'w', zipfile.ZIP_DEFLATED) as new:
        for info in z.infolist():
            data = (src / info.filename).read_bytes()
            if info.filename not in changed:
                assert data == z.read(info.filename), info.filename
            else:
                assert data != z.read(info.filename), info.filename
            new.writestr(info, data)
shutil.copy2(src / 'main.pdf', out / '237_Final.pdf')
pdf = fitz.open(out / '237_Final.pdf')
assert len(pdf) <= 20
tex = (src / 'main.tex').read_text()
assert tex.count('\\begin{figure}') == 12
assert 'fig:eval-breakdown' not in tex and 'zhangwei' not in tex
assert 'weizhang' in pdf[0].get_text()
assert '\\pagestyle{empty}' in tex and '\\thispagestyle{empty}' in tex
log = (src / 'main.log').read_text()
assert not re.search(r'undefined|multiply defined|Rerun to get', log, re.I)
pages = {}
for i, p in enumerate(pdf, 1):
    for match in re.finditer(r'^Fig\. (\d+)\.', p.get_text(), re.M):
        n = int(match.group(1))
        assert n not in pages
        pages[n] = i
assert set(pages) == set(range(1, 13))
for name in ['fig_eval_timeline_overall.pdf', 'fig_eval_timeline_groups.pdf']:
    figure = fitz.open(src / 'figures' / name)[0].get_text()
    assert 'Estimated RS visible window' in figure
    assert 'Wait/update outside forward windows' in figure
    assert 'Exposed wait' not in figure

w = openpyxl.load_workbook(old / '237_AltText.xlsx')
s = w['Alt Text']
descriptions = {
    5: 'Two averaged-cycle timelines compare the first 500 uncompressed traced cycles with the last 500 half-step compressed cycles. Rows show backward visible windows, estimated reduce-scatter (RS) visible windows, all-gather (AG) visible windows, forward hook-window unions, and recorded wait/update intervals outside forward windows, against time in milliseconds. RS windows run from host-side issue timestamps to estimated CUDA-event completion timestamps anchored to the host timeline. Forward windows are recorded by host-side hooks. Darker segments show intersections of recorded windows, not calibrated GPU-kernel overlap. The mean duration of the union of AG-wait and parameter-update intervals outside forward windows decreases from 85.0 to 35.1 milliseconds. These are instrumented visible windows, not uninstrumented throughput measurements or calibrated GPU execution intervals.',
    6: 'Per-group views of the same instrumented visible windows compare uncompressed and half-step compressed phases on a common millisecond time scale, with rows G0 through G15. Segments distinguish estimated RS visible windows, AG visible windows and group forward hook windows. Darker segments mark RS–backward and AG–forward window intersections; red segments mark recorded AG-wait and parameter-update intervals outside forward hook windows. RS completion positions are estimated from host-anchored CUDA-event timings, and forward windows are recorded by host-side hooks; the plot is not a calibrated GPU execution timeline. In the uncompressed phase, AG and forward windows extend beyond 250 milliseconds; in the compressed phase they end near 165 milliseconds. G0, which remains dense in this trace, retains a visibly longer RS window than the other compressed-phase groups.',
}
captions = {
    5: 'Instrumented visible windows before and after half-step factor pipelining. RS completion positions are estimated from host-anchored CUDA-event timings. Dark segments indicate window intersections; red segments indicate recorded AG-wait and parameter-update intervals outside forward hook windows.',
    6: 'Per-group view of the same instrumented visible windows. Each row represents one communication group; colors follow Figure 5.',
}
for row in s.iter_rows(min_row=2):
    n = int(re.search(r'Figure (\d+)', row[0].value).group(1))
    row[4].value = pages[n]
    if n in descriptions:
        row[2].value = descriptions[n]
        row[3].value = captions[n]
        s.row_dimensions[row[0].row].height = 250
info = w['Paper Information']
info['B10'] = f'Physical page index in the {len(pdf)}-page PDF; the submitted PDF has no printed page numbers.'
info['B13'] = sha(out / '237_Final.pdf')
info['B14'] = '2026-10-08'
assert info['B6'].value == 'weizhang@nudt.edu.cn'
w.save(out / '237_AltText.xlsx')
after = openpyxl.load_workbook(out / '237_AltText.xlsx')
before = openpyxl.load_workbook(old / '237_AltText.xlsx')
for a, b in zip(after['Alt Text'], before['Alt Text']):
    for x, y in zip(a, b):
        assert x._style == y._style
        if x.value != y.value:
            assert x.column == 5 or (x.row in [6, 7] and x.column in [3, 4])
shutil.copy2(old / '237_LTP.pdf', out / '237_LTP.pdf')
assert sha(out / '237_LTP.pdf') == original_hashes['237_LTP.pdf']
names = ['237_Final.pdf', '237_Source.zip', '237_LTP.pdf', '237_AltText.xlsx']
with zipfile.ZipFile(out / 'ICA3PP_237_Zhang.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    for name in names:
        z.write(out / name, name)
with zipfile.ZipFile(out / 'ICA3PP_237_Zhang.zip') as z:
    assert z.namelist() == names
    for name in names:
        assert z.read(name) == (out / name).read_bytes()
assert original_hashes == {p.name: sha(p) for p in old.iterdir() if p.is_file()}
(root / 'verification.json').write_text(json.dumps({'pages': len(pdf), 'figure_pages': pages, 'original_hashes': original_hashes}, indent=2))
for i in sorted(set([pages[5], pages[6], pages[7], pages[8], pages[9]])):
    pdf[i - 1].get_pixmap(matrix=fitz.Matrix(1.3, 1.3)).save(root / f'page_{i}.png')
print(json.dumps({'pages': len(pdf), 'figure_pages': pages, 'originals_unchanged': True}))
